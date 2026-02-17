"""
Excel Report Generator

Generates Excel reports with test results.
Columns: TestCaseID, TestName, MeasuredValue, Expected, Status
"""

from typing import List, Dict, Any
from datetime import datetime
import os


def generate_excel_report(test_results: List[Dict[str, Any]], output_dir: str = ".") -> str:
    """
    Generate an Excel report from test results.
    
    Args:
        test_results: List of test result dictionaries
        output_dir: Directory where the Excel file will be saved
    
    Returns:
        Path to the generated Excel file
    
    Raises:
        ImportError: If required Excel library is not installed
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        use_openpyxl = True
    except ImportError:
        use_openpyxl = False
    
    if not use_openpyxl:
        raise ImportError(
            "openpyxl library is required for Excel report generation.\n"
            "Install it with: pip install openpyxl"
        )
    
    # Create workbook and select active sheet
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Test Results"
    
    # Define styles
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    pass_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    fail_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    
    # Add header row
    headers = ["TestCaseID", "TestName", "MeasuredValue", "Expected", "Status"]
    worksheet.append(headers)
    
    # Style header row
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
    
    # Add data rows
    for result in test_results:
        measured_value = result.get("MeasuredValue")
        if measured_value is not None:
            measured_value = f"{measured_value:.3f}" if isinstance(measured_value, float) else str(measured_value)
        else:
            measured_value = "N/A"
        
        worksheet.append([
            result.get("TestCaseID", ""),
            result.get("TestName", ""),
            measured_value,
            result.get("Expected", ""),
            result.get("Status", "")
        ])
    
    # Style data rows
    for row_idx, result in enumerate(test_results, start=2):
        status = result.get("Status", "")
        status_fill = pass_fill if status == "PASS" else fail_fill
        status_font = Font(bold=True, color="FFFFFF")
        
        for col_idx, cell in enumerate(worksheet[row_idx], start=1):
            cell.border = thin_border
            if col_idx == len(headers):  # Status column
                cell.fill = status_fill
                cell.font = status_font
                cell.alignment = center_alignment
            elif col_idx == 1:  # TestCaseID column
                cell.alignment = left_alignment
            else:
                cell.alignment = left_alignment
    
    # Adjust column widths
    worksheet.column_dimensions["A"].width = 15
    worksheet.column_dimensions["B"].width = 35
    worksheet.column_dimensions["C"].width = 15
    worksheet.column_dimensions["D"].width = 20
    worksheet.column_dimensions["E"].width = 12
    
    # Add summary sheet
    summary_sheet = workbook.create_sheet("Summary")
    total_tests = len(test_results)
    passed_tests = sum(1 for r in test_results if r.get("Status") == "PASS")
    failed_tests = sum(1 for r in test_results if r.get("Status") == "FAIL")
    
    summary_sheet["A1"] = "Test Execution Summary"
    summary_sheet["A1"].font = Font(bold=True, size=14)
    
    summary_sheet["A3"] = "Total Tests:"
    summary_sheet["B3"] = total_tests
    
    summary_sheet["A4"] = "Passed:"
    summary_sheet["B4"] = passed_tests
    summary_sheet["B4"].fill = pass_fill
    
    summary_sheet["A5"] = "Failed:"
    summary_sheet["B5"] = failed_tests
    if failed_tests > 0:
        summary_sheet["B5"].fill = fail_fill
    
    summary_sheet.column_dimensions["A"].width = 20
    summary_sheet.column_dimensions["B"].width = 15
    
    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"SmartBU_TestResults_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)
    
    # Ensure output directory exists
    os.makedirs(output_dir, exist_ok=True)
    
    # Save workbook
    workbook.save(filepath)
    
    return filepath
