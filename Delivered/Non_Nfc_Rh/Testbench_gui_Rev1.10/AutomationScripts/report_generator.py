"""
Test Report Generator
=====================
Copies the Excel test-specification template and fills in the **Observed Result**
and **Status** columns using the actual hardware measurements produced by
:meth:`~AutomationScripts.core.test_sequences.TestSequenceRunner.run_for_variant`.

Typical usage (called automatically from ``IntegratedAutomationRunner``)::

    from AutomationScripts.report_generator import generate_report
    path = generate_report(run_results=results)   # results = run_for_variant()
    print(f"Report saved to: {path}")

NFC, CAN and LIN sheets are left untouched (not applicable for this variant).
"""

from __future__ import annotations

import os
import shutil
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

# ── Paths ─────────────────────────────────────────────────────────────────────
_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.join(_HERE, "..")

TEMPLATE_PATH = os.path.join(_ROOT, "AutomationTest", "Smart_BU_Test Specification.xlsx")
REPORTS_DIR   = os.path.join(_ROOT, "AutomationTest", "reports")

# ── Cell fill colours ─────────────────────────────────────────────────────────
_PASS_FILL = PatternFill("solid", fgColor="92D050")   # green
_FAIL_FILL = PatternFill("solid", fgColor="FF4C4C")   # red
_SKIP_FILL = PatternFill("solid", fgColor="FFFF00")   # yellow (not run)

# ── Column indices (1-based) in the template ──────────────────────────────────
_COL_TC_ID     = 2   # B  Test Case ID
_COL_PRE_ACT   = 4   # D  Pre Action
_COL_STEPS     = 5   # E  Test Steps
_COL_EXPECTED  = 6   # F  Expected Result
_COL_POST_ACT  = 7   # G  Post Action
_COL_OBS       = 8   # H  Observed Result
_COL_STATUS    = 9   # I  Status

_WRAP_TOP = Alignment(wrap_text=True, vertical="top")

# ── Shared automation pre/post actions ────────────────────────────────────────
_PRE_COMMON = (
    "1. Power Supply ON (12 V)\n"
    "2. Trace32 hardware connected to PC and target board\n"
    "3. ELF selected and flashed via automation\n"
    "4. Trace32 connected and firmware running (automated)"
)
_POST_COMMON = "Automated cleanup — Power Supply turned OFF after run"


# ── Convert run_for_variant() output → per-sheet row dicts ────────────────────

def _row(
    tc_id: str,
    tc_name: str,
    pre_action: str,
    test_steps: str,
    expected: str,
    observed: str,
    status: str,
    post_action: str = _POST_COMMON,
) -> Dict[str, Any]:
    return {
        "TestCaseID":   tc_id,
        "TestName":     tc_name,
        "PreAction":    pre_action,
        "TestSteps":    test_steps,
        "Expected":     expected,
        "ObservedText": observed,
        "Status":       status,
        "PostAction":   post_action,
    }


def _status(passed) -> str:
    if isinstance(passed, dict):
        passed = passed.get("pass", False)
    return "PASS" if passed else "FAIL"


def results_from_run(run_results: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    """
    Convert the dictionary returned by ``TestSequenceRunner.run_for_variant()``
    into a per-sheet list of row dicts ready for :func:`_fill_sheet`.

    The run_results keys and their types::

        battery  : {'pass': bool, 'voltage': float}
        capa1    : {'pass': bool, 'TestFw_CapaUnlockSensorValue': float, ...}
        capa2    : same structure as capa1
        led_on   : {'pass': bool, 'voltage': float}
        led_off  : {'pass': bool, 'voltage': float}
        motor    : {'pass': bool, 'voltage': float, 'current': float, 'load_error': float}
        eos_reset: {'pass': bool, 'voltage': float}
        eos_set  : {'pass': bool, 'voltage': float}
        sg1      : {'pass': bool, 'plus': float, 'minus': float, 'opamp': float}
        sg2      : {'pass': bool, 'plus': float, 'minus': float, 'opamp': float}
    """
    by_sheet: Dict[str, List[Dict[str, Any]]] = {}

    # ── Battery ───────────────────────────────────────────────────────────────
    bat = run_results.get("battery", {})
    if isinstance(bat, dict):
        bat_v = bat.get("voltage", 0.0)
        by_sheet["Battery"] = [_row(
            "TC_BAT_01",
            "Test case to verify Battery monitor of the Non-Driver PCB",
            pre_action=_PRE_COMMON,
            test_steps=(
                "1. Automation sends DID: TESTFW_GUI_CMD_BATT_MONITOR_e\n"
                "2. Waits 2 s for ADC stabilisation\n"
                "3. Polls TestFw_AiBatRef until two consecutive identical readings\n"
                "4. If first reading is 0.0 mV, waits 10 s and retries once"
            ),
            expected="AiBatRef = 11500-13500 mV",
            observed=f"AiBatRef = {bat_v:.1f} mV",
            status=_status(bat),
        )]

    # ── LED ───────────────────────────────────────────────────────────────────
    led_rows: List[Dict[str, Any]] = []
    led_on = run_results.get("led_on", {})
    if isinstance(led_on, dict):
        v = led_on.get("voltage", 0.0)
        led_rows.append(_row(
            "TC_LED_01",
            "Test case to verify LED voltage of the Non-Driver PCB when LED connected",
            pre_action=(
                _PRE_COMMON + "\n"
                "5. LED connected between TP602 and GND"
            ),
            test_steps=(
                "1. Set LedTest_LedCanLinRequest = 1 (LED ON)\n"
                "2. Wait 2 s for voltage to stabilise\n"
                "3. Send DID: TESTFW_GUI_CMD_LED_TEST_e\n"
                "4. Poll TestFw_LedVoltage until stable"
            ),
            expected="LedVoltage = 2400-2600 mV",
            observed=f"LedVoltage = {v:.1f} mV",
            status=_status(led_on),
        ))
    led_off = run_results.get("led_off", {})
    if isinstance(led_off, dict):
        v = led_off.get("voltage", 0.0)
        led_rows.append(_row(
            "TC_LED_03",
            "Test case to verify LED voltage of the Non-Driver PCB when LED not connected",
            pre_action=(
                _PRE_COMMON + "\n"
                "5. LED not connected (TP602 open / no current path)"
            ),
            test_steps=(
                "1. Set LedTest_LedCanLinRequest = 0 (LED OFF)\n"
                "2. Wait 2 s for voltage to stabilise\n"
                "3. Send DID: TESTFW_GUI_CMD_LED_TEST_e\n"
                "4. Poll TestFw_LedVoltage until stable"
            ),
            expected="LedVoltage <= 10 mV",
            observed=f"LedVoltage = {v:.1f} mV",
            status=_status(led_off),
        ))
    if led_rows:
        by_sheet["LED "] = led_rows

    # ── Motor ─────────────────────────────────────────────────────────────────
    mot = run_results.get("motor", {})
    if isinstance(mot, dict):
        v = mot.get("voltage", 0.0)
        i = mot.get("current", 0.0)
        e = mot.get("load_error", -1.0)
        by_sheet["Motor"] = [_row(
            "TC_MOTOR_01",
            "Test case to verify motor functionality of the Non-Driver PCB",
            pre_action=(
                _PRE_COMMON + "\n"
                "5. Actuator connected between DRV_OUT_A and DRV_OUT_B motor driver terminals"
            ),
            test_steps=(
                "1. Set MotorTest_SetGuiMotorActuateRequest = 1 (actuate)\n"
                "2. Wait 1 s, then clear request to 0\n"
                "3. Send DID: TESTFW_GUI_CMD_MOTOR_TEST_e\n"
                "4. Poll TestFw_MotorVoltage, TestFw_MotorCurrentValue, TestFw_MotorLoadError until stable"
            ),
            expected=(
                "MotorVoltage > 0 mV\n"
                "MotorCurrent > 0 mA (and != 65535 mA)\n"
                "MotorLoadError = 0"
            ),
            observed=(
                f"MotorVoltage = {v:.1f} mV\n"
                f"MotorCurrent = {i:.1f} mA\n"
                f"MotorLoadError = {int(e)}"
            ),
            status=_status(mot),
        )]

    # ── EOS ───────────────────────────────────────────────────────────────────
    eos_rows: List[Dict[str, Any]] = []
    _eos_pre = (
        _PRE_COMMON + "\n"
        "5. 560 Ohm pull-up resistor connected between EOS_OUT and 3.3 V supply\n"
        "6. IPB_GND shorted to external 3.3 V supply ground"
    )
    eos_rst = run_results.get("eos_reset", {})
    if isinstance(eos_rst, dict):
        v = eos_rst.get("voltage", 0.0)
        eos_rows.append(_row(
            "TC_EOS_01",
            "Test case to verify EOS functionality of the Non-Driver PCB (Reset state)",
            pre_action=_eos_pre,
            test_steps=(
                "1. Set EosTest_EosRequestGui = 0 (Reset / DO_EOS NOT ACTIVE)\n"
                "2. Send DID: TESTFW_GUI_CMD_EOS_TEST_e (reset command)\n"
                "3. Wait 1 s for stabilisation\n"
                "4. Poll TestFw_EosDiagVoltage until stable"
            ),
            expected="EosDiagVoltage = 1500-3000 mV",
            observed=f"EosDiagVoltage = {v:.1f} mV",
            status=_status(eos_rst),
        ))
    eos_set = run_results.get("eos_set", {})
    if isinstance(eos_set, dict):
        v = eos_set.get("voltage", 0.0)
        eos_rows.append(_row(
            "TC_EOS_02",
            "Test case to verify EOS functionality of the Non-Driver PCB (Set state)",
            pre_action=_eos_pre,
            test_steps=(
                "1. Clear entry field, wait 0.5 s\n"
                "2. Set EosTest_EosRequestGui = 1 (Set / DO_EOS ACTIVE)\n"
                "3. Send DID: TESTFW_GUI_CMD_EOS_TEST_e (set command)\n"
                "4. Wait 2 s for set process to complete\n"
                "5. Poll TestFw_EosDiagVoltage until stable"
            ),
            expected="EosDiagVoltage = 1400-1600 mV",
            observed=f"EosDiagVoltage = {v:.1f} mV",
            status=_status(eos_set),
        ))
    if eos_rows:
        by_sheet["EOS"] = eos_rows

    # ── SG ────────────────────────────────────────────────────────────────────
    sg_rows: List[Dict[str, Any]] = []
    _sg_pre = (
        _PRE_COMMON + "\n"
        "5. Dummy Wheatstone bridge (4 x 1.2 kOhm) connected to SG1+/SG1- and SG2+/SG2-"
    )
    sg1 = run_results.get("sg1", {})
    if isinstance(sg1, dict):
        p = sg1.get("plus", 0.0); m = sg1.get("minus", 0.0); o = sg1.get("opamp", 0.0)
        sg_rows.append(_row(
            "TC_SG_01",
            "Test case for SG1 Operational Amplifier Verification",
            pre_action=_sg_pre,
            test_steps=(
                "1. Clear SG entry field\n"
                "2. Set TestFw_GetSgResults = 1\n"
                "3. Send DID: TEST_GUI_CMD_SG_TEST_e\n"
                "4. Wait 3 s for measurement\n"
                "5. Re-send DID: TEST_GUI_CMD_SG_TEST_e\n"
                "6. Poll TestFw_Sg1PlusOpamp, TestFw_Sg1MinusOpamp, TestFw_Sg1Opamp until stable"
            ),
            expected=(
                "SG1+ (TestFw_Sg1PlusOpamp) = 1600-4100 mV\n"
                "SG1- (TestFw_Sg1MinusOpamp) = 1600-4100 mV\n"
                "SG1_OpAmp (TestFw_Sg1Opamp) = 200-4600 mV"
            ),
            observed=(
                f"SG1+ = {p:.1f} mV\n"
                f"SG1- = {m:.1f} mV\n"
                f"SG1_OpAmp = {o:.1f} mV"
            ),
            status=_status(sg1),
        ))
    sg2 = run_results.get("sg2", {})
    if isinstance(sg2, dict):
        p = sg2.get("plus", 0.0); m = sg2.get("minus", 0.0); o = sg2.get("opamp", 0.0)
        sg_rows.append(_row(
            "TC_SG_03",
            "Test case for SG2 Operational Amplifier Verification",
            pre_action=_sg_pre,
            test_steps=(
                "1. Reset TestFw_GetSgResults = 0, wait 1 s\n"
                "2. Clear SG entry field, wait 1 s\n"
                "3. Set TestFw_GetSgResults = 1\n"
                "4. Send DID: TEST_GUI_CMD_SG_TEST_e\n"
                "5. Wait 3 s for measurement\n"
                "6. Re-send DID: TEST_GUI_CMD_SG_TEST_e\n"
                "7. Poll TestFw_Sg2PlusOpamp, TestFw_Sg2MinusOpamp, TestFw_Sg2Opamp until stable"
            ),
            expected=(
                "SG2+ (TestFw_Sg2PlusOpamp) = 1600-4100 mV\n"
                "SG2- (TestFw_Sg2MinusOpamp) = 1600-4100 mV\n"
                "SG2_OpAmp (TestFw_Sg2Opamp) = 200-4600 mV"
            ),
            observed=(
                f"SG2+ = {p:.1f} mV\n"
                f"SG2- = {m:.1f} mV\n"
                f"SG2_OpAmp = {o:.1f} mV"
            ),
            status=_status(sg2),
        ))
    if sg_rows:
        by_sheet["SG"] = sg_rows

    # ── CAPA ──────────────────────────────────────────────────────────────────
    capa_rows: List[Dict[str, Any]] = []
    _capa_pre = (
        _PRE_COMMON + "\n"
        "5. 100 pF capacitor connected to Unlock, Approach and Lock sensor channels via switch"
    )
    for tc_id, key, label, sensor_key, sensor_name, threshold in [
        ("TC_CAPA_01", "capa1", "1st", "TestFw_CapaUnlockSensorValue",   "Unlock",   "> 8900"),
        ("TC_CAPA_02", "capa2", "2nd", "TestFw_CapaApproachSensorValue", "Approach", "> 9000"),
        ("TC_CAPA_03", "capa1", "1st", "TestFw_CapaLockSensorValue",     "Lock",     "> 9000"),
    ]:
        c = run_results.get(key, {})
        if isinstance(c, dict):
            u = c.get("TestFw_CapaUnlockSensorValue",   0.0)
            a = c.get("TestFw_CapaApproachSensorValue", 0.0)
            l = c.get("TestFw_CapaLockSensorValue",     0.0)
            sensor_val = c.get(sensor_key, 0.0)
            sensor_passed = isinstance(sensor_val, (int, float)) and sensor_val > (8900 if "Unlock" in sensor_name else 9000)
            capa_rows.append(_row(
                tc_id,
                f"Test case to verify CAPA {sensor_name} sensor functionality",
                pre_action=_capa_pre,
                test_steps=(
                    f"1. Automation sends DID: TEST_GUI_CMD_CAPA_TEST_e ({label} attempt)\n"
                    "2. Retries every 0.5 s for up to 10 s until all sensor values are in active range\n"
                    f"3. Reads TestFw_Capa{sensor_name}SensorValue\n"
                    f"4. Reads TestFw_Capa{sensor_name} flag"
                ),
                expected=f"TestFw_Capa{sensor_name}SensorValue {threshold}",
                observed=(
                    f"UnlockSensor = {u:.0f}\n"
                    f"ApproachSensor = {a:.0f}\n"
                    f"LockSensor = {l:.0f}"
                ),
                status="PASS" if sensor_passed else "FAIL",
            ))
    if capa_rows:
        by_sheet["Capa"] = capa_rows

    return by_sheet


# ── Worksheet filling ─────────────────────────────────────────────────────────

def _fill_sheet(ws, results: List[Dict[str, Any]]) -> None:
    """Overwrite Pre Action, Test Steps, Expected Result, Post Action,
    Observed Result and Status for every matched test-case row."""
    by_tc_id: Dict[str, Dict[str, Any]] = {
        str(r.get("TestCaseID", "")).strip(): r for r in results
        if r.get("TestCaseID")
    }

    for row_idx in range(2, ws.max_row + 1):
        tc_id = str(ws.cell(row=row_idx, column=_COL_TC_ID).value or "").strip()
        result = by_tc_id.get(tc_id)
        if not tc_id or result is None:
            continue

        # ── Pre Action (D) ────────────────────────────────────────────────────
        c = ws.cell(row=row_idx, column=_COL_PRE_ACT)
        c.value     = result.get("PreAction", "")
        c.alignment = _WRAP_TOP

        # ── Test Steps (E) ────────────────────────────────────────────────────
        c = ws.cell(row=row_idx, column=_COL_STEPS)
        c.value     = result.get("TestSteps", "")
        c.alignment = _WRAP_TOP

        # ── Expected Result (F) ───────────────────────────────────────────────
        c = ws.cell(row=row_idx, column=_COL_EXPECTED)
        c.value     = result.get("Expected", "")
        c.alignment = _WRAP_TOP

        # ── Post Action (G) ───────────────────────────────────────────────────
        c = ws.cell(row=row_idx, column=_COL_POST_ACT)
        c.value     = result.get("PostAction", "")
        c.alignment = _WRAP_TOP

        # ── Observed Result (H) ───────────────────────────────────────────────
        c = ws.cell(row=row_idx, column=_COL_OBS)
        c.value     = result.get("ObservedText", "")
        c.alignment = _WRAP_TOP

        # ── Status (I) ────────────────────────────────────────────────────────
        status    = result.get("Status", "")
        stat_cell = ws.cell(row=row_idx, column=_COL_STATUS)
        stat_cell.value     = status
        stat_cell.font      = Font(bold=True)
        stat_cell.alignment = Alignment(horizontal="center", vertical="center")
        stat_cell.fill = (
            _PASS_FILL if status == "PASS" else
            _FAIL_FILL if status == "FAIL" else
            _SKIP_FILL
        )


# ── Public API ────────────────────────────────────────────────────────────────

def generate_report(
    run_results: Optional[Dict[str, Any]] = None,
    output_path: Optional[str] = None,
) -> str:
    """
    Generate a filled-in Excel test report from the specification template.

    Parameters
    ----------
    run_results:
        The dictionary returned by
        ``TestSequenceRunner.run_for_variant()``.  When *None* an empty
        report (template copy) is produced.
    output_path:
        Destination file path.  Defaults to
        ``AutomationTest/reports/<YYYYMMDD_HHMMSS>_Test_Report.xlsx``.

    Returns
    -------
    str
        Absolute path to the generated report file.
    """
    results_by_sheet = results_from_run(run_results) if run_results else {}

    os.makedirs(REPORTS_DIR, exist_ok=True)

    if output_path is None:
        ts          = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(REPORTS_DIR, f"{ts}_Test_Report.xlsx")

    output_path = os.path.abspath(output_path)
    shutil.copy2(os.path.abspath(TEMPLATE_PATH), output_path)

    wb = openpyxl.load_workbook(output_path)

    # Remove sheets not relevant to automated testing
    for sheet_to_remove in ("Voltage_Check", "Common Operations_Preconditions"):
        if sheet_to_remove in wb.sheetnames:
            del wb[sheet_to_remove]

    for sheet_name, rows in results_by_sheet.items():
        if sheet_name in wb.sheetnames and rows:
            _fill_sheet(wb[sheet_name], rows)
            # Delete any template row whose TC_ID was not executed
            executed_ids = {
                str(r.get("TestCaseID", "")).strip()
                for r in rows if r.get("TestCaseID")
            }
            ws = wb[sheet_name]
            to_delete = [
                row_idx
                for row_idx in range(2, ws.max_row + 1)
                if str(ws.cell(row=row_idx, column=_COL_TC_ID).value or "").strip() not in ("", *executed_ids)
            ]
            for row_idx in reversed(to_delete):
                ws.delete_rows(row_idx)

    wb.save(output_path)
    wb.close()
    return output_path
